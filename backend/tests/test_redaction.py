"""Unit tests for document redaction — RedactionService and redact_document task."""


# ---------------------------------------------------------------------------
# Helpers — build minimal in-memory fixtures
# ---------------------------------------------------------------------------


def _make_pdf(text: str, author: str = "") -> bytes:
    import io

    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 100), text)
    if author:
        doc.set_metadata({"author": author})
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def _make_docx(text: str, author: str = "") -> bytes:
    import io

    from docx import Document

    doc = Document()
    if author:
        doc.core_properties.author = author
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# RedactionService.build_keyword_map
# ---------------------------------------------------------------------------


def test_build_keyword_map_assigns_stable_placeholders():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    km = svc.build_keyword_map(frozenset({"Alice", "Bob"}))
    assert km["Alice"] == "[REDACTED_1]"
    assert km["Bob"] == "[REDACTED_2]"


def test_build_keyword_map_deduplicates_case_insensitively():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    km = svc.build_keyword_map(frozenset({"John", "john"}))
    assert km["John"] == km["john"]
    assert len(set(km.values())) == 1


def test_build_keyword_map_empty_returns_empty():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    assert svc.build_keyword_map(frozenset()) == {}


# ---------------------------------------------------------------------------
# RedactionService.redact_text
# ---------------------------------------------------------------------------


def test_redact_text_replaces_keyword():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    text, counts = svc.redact_text(
        "Hello John, how are you John?", {"John": "[REDACTED_1]"}
    )
    assert "[REDACTED_1]" in text
    assert "John" not in text
    assert counts == {"[REDACTED_1]": 2}


def test_redact_text_is_case_insensitive():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    text, counts = svc.redact_text("Hello John and JOHN", {"john": "[REDACTED_1]"})
    assert "John" not in text
    assert "JOHN" not in text
    assert counts["[REDACTED_1]"] == 2


def test_redact_text_word_boundary_prevents_partial_match():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    text, counts = svc.redact_text("Jonathan Jones Jon", {"Jon": "[REDACTED_1]"})
    assert "Jonathan" in text
    assert counts.get("[REDACTED_1]", 0) == 1


def test_redact_text_longest_first_prevents_shadowing():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    km = {"John Smith": "[REDACTED_1]", "John": "[REDACTED_2]"}
    text, counts = svc.redact_text("John Smith is here", km)
    assert "[REDACTED_1]" in text
    assert "John Smith" not in text
    assert "[REDACTED_2]" not in text


def test_redact_text_empty_keyword_map_is_no_op():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    original = "Nothing changes here"
    text, counts = svc.redact_text(original, {})
    assert text == original
    assert counts == {}


def test_redact_text_returns_occurrence_counts_per_placeholder():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    km = {"Alice": "[REDACTED_1]", "Bob": "[REDACTED_2]"}
    _, counts = svc.redact_text("Alice met Bob and Alice left", km)
    assert counts["[REDACTED_1]"] == 2
    assert counts["[REDACTED_2]"] == 1


# ---------------------------------------------------------------------------
# RedactionService.redact_pdf
# ---------------------------------------------------------------------------


def test_redact_pdf_returns_valid_pdf_bytes():
    import fitz

    from app.redaction.service import RedactionService

    svc = RedactionService()
    result_bytes, _, _ = svc.redact_pdf(_make_pdf("Hello world"), {})
    doc = fitz.open(stream=result_bytes, filetype="pdf")
    assert doc.page_count == 1
    doc.close()


def test_redact_pdf_scrubs_metadata():
    import fitz

    from app.redaction.service import RedactionService

    svc = RedactionService()
    result_bytes, _, orig_metadata = svc.redact_pdf(
        _make_pdf("Text", author="Jane Doe"), {}
    )
    assert orig_metadata.get("author") == "Jane Doe"

    out_doc = fitz.open(stream=result_bytes, filetype="pdf")
    assert not out_doc.metadata.get("author")
    out_doc.close()


def test_redact_pdf_counts_keyword_occurrences():
    from app.redaction.service import RedactionService

    svc = RedactionService()
    pdf = _make_pdf("John Smith saw John leave.")
    _, counts, _ = svc.redact_pdf(pdf, {"John": "[REDACTED_1]"})
    assert counts.get("[REDACTED_1]", 0) >= 1


def test_redact_pdf_empty_keyword_map_is_no_op():
    import fitz

    from app.redaction.service import RedactionService

    svc = RedactionService()
    result_bytes, counts, _ = svc.redact_pdf(_make_pdf("Untouched text"), {})
    assert counts == {}
    doc = fitz.open(stream=result_bytes, filetype="pdf")
    assert doc.page_count == 1
    doc.close()


# ---------------------------------------------------------------------------
# RedactionService.redact_docx
# ---------------------------------------------------------------------------


def test_redact_docx_replaces_keyword_in_paragraph():
    import io

    from docx import Document

    from app.redaction.service import RedactionService

    svc = RedactionService()
    result_bytes, counts, _ = svc.redact_docx(
        _make_docx("John Smith was present."), {"John": "[REDACTED_1]"}
    )
    out_text = " ".join(p.text for p in Document(io.BytesIO(result_bytes)).paragraphs)
    assert "[REDACTED_1]" in out_text
    assert counts.get("[REDACTED_1]", 0) >= 1


def test_redact_docx_scrubs_author_metadata():
    import io

    from docx import Document

    from app.redaction.service import RedactionService

    svc = RedactionService()
    result_bytes, _, orig_metadata = svc.redact_docx(
        _make_docx("Content", author="Jane Doe"), {}
    )
    assert orig_metadata.get("author") == "Jane Doe"
    assert not Document(io.BytesIO(result_bytes)).core_properties.author


def test_redact_docx_empty_keyword_map_is_no_op():
    import io

    from docx import Document

    from app.redaction.service import RedactionService

    svc = RedactionService()
    result_bytes, counts, _ = svc.redact_docx(_make_docx("Untouched text"), {})
    assert counts == {}
    out_text = " ".join(p.text for p in Document(io.BytesIO(result_bytes)).paragraphs)
    assert "Untouched text" in out_text


# ---------------------------------------------------------------------------
# redact_document Celery task
# ---------------------------------------------------------------------------


def test_redact_document_task_is_registered():
    import app.workers.tasks.redact_document  # noqa: F401
    from app.workers import celery_app

    assert "gideon.redact_document" in celery_app.tasks


def test_redact_document_in_task_registry():
    from app.workers.registry import TASK_REGISTRY

    assert TASK_REGISTRY["redact_document"] == "gideon.redact_document"


def test_redact_document_skips_when_no_keywords_configured():
    from unittest.mock import MagicMock, patch

    from app.workers.tasks.redact_document import redact_document

    mock_settings = MagicMock()
    mock_settings.redaction.keywords = frozenset()

    with patch("app.workers.tasks.redact_document.settings", mock_settings):
        result = redact_document("doc-1", "firm/matter/doc-1/original.pdf")

    assert result["status"] == "skipped"
    assert result["reason"] == "no_keywords_configured"


def test_redact_document_skips_unsupported_format():
    from unittest.mock import AsyncMock, MagicMock, patch

    from app.workers.tasks.redact_document import redact_document

    mock_settings = MagicMock()
    mock_settings.redaction.keywords = frozenset({"John"})

    mock_storage = AsyncMock()
    mock_storage.download_document.return_value = (b"plain data", "text/plain")

    mock_orig = {
        "firm_id": "firm-1",
        "matter_id": "matter-1",
        "filename": "report.txt",
    }

    with (
        patch("app.workers.tasks.redact_document.settings", mock_settings),
        patch(
            "app.workers.tasks.redact_document._load_original_document",
            new=AsyncMock(return_value=mock_orig),
        ),
        patch("app.storage.get_storage_service", return_value=mock_storage),
    ):
        result = redact_document("doc-1", "firm/matter/doc-1/report.txt")

    assert result["status"] == "skipped"
    assert result["reason"] == "unsupported_format"


def test_redact_document_full_pipeline_pdf():
    from unittest.mock import AsyncMock, MagicMock, patch

    from app.redaction.service import RedactionService
    from app.workers.tasks.redact_document import redact_document

    pdf_bytes = _make_pdf("John Smith is a witness.")

    mock_settings = MagicMock()
    mock_settings.redaction.keywords = frozenset({"John"})

    mock_storage = AsyncMock()
    mock_storage.download_document.return_value = (pdf_bytes, "application/pdf")
    mock_storage.upload_bytes = AsyncMock()
    mock_storage.upload_json = AsyncMock()

    mock_orig = {
        "firm_id": "firm-1",
        "matter_id": "matter-1",
        "filename": "original.pdf",
    }

    with (
        patch("app.workers.tasks.redact_document.settings", mock_settings),
        patch(
            "app.workers.tasks.redact_document._load_original_document",
            new=AsyncMock(return_value=mock_orig),
        ),
        patch(
            "app.workers.tasks.redact_document._update_has_redacted_copy",
            new=AsyncMock(),
        ),
        patch("app.storage.get_storage_service", return_value=mock_storage),
        patch("app.redaction.get_redaction_service", return_value=RedactionService()),
        patch("app.workers.celery_app.send_task") as mock_send,
    ):
        result = redact_document("doc-1", "firm/matter/doc-1/original.pdf")

    assert result["status"] == "completed"
    assert result["original_document_id"] == "doc-1"
    assert result["keyword_count"] == 1
    assert "redacted_document_id" not in result
    assert result["redacted_s3_key"] == "firm/matter/doc-1/redacted.pdf"
    mock_storage.upload_bytes.assert_awaited_once()
    mock_storage.upload_json.assert_awaited_once()
    # ingest is re-dispatched for the SAME document with the sidecar key
    mock_send.assert_called_once()
    call_args = mock_send.call_args
    assert call_args[0][0] == "gideon.ingest_document"
    assert call_args.kwargs["args"][0] == "doc-1"
    assert call_args.kwargs["args"][1] == "firm/matter/doc-1/redacted.pdf"


def test_redact_document_full_pipeline_docx():
    from unittest.mock import AsyncMock, MagicMock, patch

    from app.redaction.service import RedactionService
    from app.workers.tasks.redact_document import redact_document

    docx_bytes = _make_docx("Jane Doe filed the motion.")
    docx_mime = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    mock_settings = MagicMock()
    mock_settings.redaction.keywords = frozenset({"Jane Doe"})

    mock_storage = AsyncMock()
    mock_storage.download_document.return_value = (docx_bytes, docx_mime)
    mock_storage.upload_bytes = AsyncMock()
    mock_storage.upload_json = AsyncMock()

    mock_orig = {
        "firm_id": "firm-1",
        "matter_id": "matter-1",
        "filename": "motion.docx",
    }

    with (
        patch("app.workers.tasks.redact_document.settings", mock_settings),
        patch(
            "app.workers.tasks.redact_document._load_original_document",
            new=AsyncMock(return_value=mock_orig),
        ),
        patch(
            "app.workers.tasks.redact_document._update_has_redacted_copy",
            new=AsyncMock(),
        ),
        patch("app.storage.get_storage_service", return_value=mock_storage),
        patch("app.redaction.get_redaction_service", return_value=RedactionService()),
        patch("app.workers.celery_app.send_task") as mock_send,
    ):
        result = redact_document("doc-2", "firm/matter/doc-2/motion.docx")

    assert result["status"] == "completed"
    assert result["original_document_id"] == "doc-2"
    assert "redacted_document_id" not in result
    assert result["redacted_s3_key"] == "firm/matter/doc-2/redacted.docx"
    assert result["keyword_count"] == 1
    assert result["total_replacements"] >= 1
    call_args = mock_send.call_args
    assert call_args[0][0] == "gideon.ingest_document"
    assert call_args.kwargs["args"][0] == "doc-2"
    assert call_args.kwargs["args"][1] == "firm/matter/doc-2/redacted.docx"
